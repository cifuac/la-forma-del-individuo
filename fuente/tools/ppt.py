"""Versión en PowerPoint: cada lámina como imagen a pantalla completa, con su título (detrás de la imagen),
sus notas de orador y las secciones de la presentación.

Uso: python3 -I tools/ppt.py <raiz_web> <dir_capturas> <salida.pptx>
"""
import html
import os
import re
import sys
import uuid
from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches

raiz, capturas, salida = sys.argv[1:4]
fuente = open(os.path.join(raiz, 'dist', 'local.html'), encoding='utf-8').read()
laminas = {}
for m in re.finditer(r'<section class="lamina[^"]*" id="([^"]+)"([^>]*)>(.*?)</section>', fuente, re.S):
    attrs, cuerpo = m.group(2), m.group(3)
    titulo = re.search(r'data-titulo="([^"]*)"', attrs)
    seccion = re.search(r'data-seccion="([^"]*)"', attrs)
    notas = re.search(r'<aside class="notas">(.*?)</aside>', cuerpo, re.S)
    texto = html.unescape(re.sub(r'<[^>]+>', '', notas.group(1))).strip() if notas else ''
    laminas[m.group(1)] = (html.unescape(titulo.group(1)) if titulo else m.group(1), html.unescape(seccion.group(1)) if seccion else '', texto)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
solo_titulo = prs.slide_layouts[5]
secciones = []
for nombre in sorted(f for f in os.listdir(capturas) if f.endswith(('.png', '.jpg'))):
    m = re.match(r'\d+_(.+?)(?:_p(\d+))?\.(?:png|jpg)$', nombre)
    ident, paso = m.group(1), m.group(2)
    titulo, seccion, notas = laminas[ident]
    if paso is not None:
        total = len([f for f in os.listdir(capturas) if re.match(r'\d+_%s_p\d+\.(?:png|jpg)$' % re.escape(ident), f)])
        titulo = '%s (%d de %d)' % (titulo, int(paso) + 1, total)
    s = prs.slides.add_slide(solo_titulo)
    s.shapes.title.text = titulo
    foto = s.shapes.add_picture(os.path.join(capturas, nombre), 0, 0, width=prs.slide_width, height=prs.slide_height)
    foto.name = titulo
    foto._element.nvPicPr.cNvPr.set('descr', titulo)
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    if not secciones or secciones[-1][0] != seccion:
        secciones.append((seccion, []))
    secciones[-1][1].append(s.slide_id)

# secciones (PowerPoint 2010+): una por parte de la presentación
P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'
NS_P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
pres = prs.part._element
ext_lst = pres.find('{%s}extLst' % NS_P)
if ext_lst is None:
    ext_lst = etree.SubElement(pres, '{%s}extLst' % NS_P)
ext = etree.SubElement(ext_lst, '{%s}ext' % NS_P, uri='{521415D9-36F7-43E2-AB2F-B90AF26B5E84}')
lista = etree.SubElement(ext, '{%s}sectionLst' % P14, nsmap={'p14': P14})
for nombre, ids in secciones:
    sec = etree.SubElement(lista, '{%s}section' % P14, name=nombre or 'Presentación', id='{%s}' % str(uuid.uuid4()).upper())
    sl = etree.SubElement(sec, '{%s}sldIdLst' % P14)
    for i in ids:
        etree.SubElement(sl, '{%s}sldId' % P14, id=str(i))

prs.core_properties.title = 'La forma del individuo · Ian Watt y el surgimiento de la novela'
prs.save(salida)
print('diapositivas:', len(prs.slides), '· secciones:', len(secciones), '·', round(os.path.getsize(salida) / 1e6, 1), 'MB')
